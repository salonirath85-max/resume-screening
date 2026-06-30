import os
from typing import Dict

import pdfplumber
from docx import Document
from fpdf import FPDF


def extract_text_from_file(file_path: str) -> str:
    """Extract readable text from PDF, DOCX, or TXT files."""
    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as handle:
            return handle.read()

    if extension == ".pdf":
        try:
            with pdfplumber.open(file_path) as pdf:
                pages = [page.extract_text() or "" for page in pdf.pages]
                return "\n".join(pages)
        except Exception as exc:  # pragma: no cover - defensive fallback
            raise RuntimeError(f"Unable to read PDF: {exc}") from exc

    if extension == ".docx":
        try:
            document = Document(file_path)
            return "\n".join(paragraph.text for paragraph in document.paragraphs)
        except Exception as exc:  # pragma: no cover - defensive fallback
            raise RuntimeError(f"Unable to read DOCX: {exc}") from exc

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX, or TXT file.")


def create_analysis_pdf(analysis: Dict[str, object]) -> bytes:
    """Create a simple PDF report for the analysis results."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(200, 10, "Resume Screening Report", ln=True, align="C")
    pdf.ln(5)

    pdf.set_font("Arial", size=11)
    pdf.multi_cell(0, 8, f"Match Score: {analysis['score']:.1f}%")
    pdf.multi_cell(0, 8, f"Rating: {analysis['rating_label']} ({analysis['rating_text']})")
    pdf.multi_cell(0, 8, f"Word Count: {analysis['word_count']}")
    pdf.multi_cell(0, 8, f"Completeness: {analysis['completeness_score']:.1f}%")
    pdf.ln(3)
    pdf.multi_cell(0, 8, "Skills Found: " + ", ".join(analysis["matched_skills"]))
    pdf.multi_cell(0, 8, "Missing Skills: " + ", ".join(analysis["missing_skills"]))
    pdf.multi_cell(0, 8, "Suggestions: " + " | ".join(analysis["suggestions"]))

    output = pdf.output(dest="S")
    return output
